"""Turn a raw orders CSV into a formatted Excel sales report.
Usage: python sales_report.py orders.csv report.xlsx
"""
import sys
import pandas as pd
from openpyxl.chart import BarChart, LineChart, Reference
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter

def build(src, out):
    df = pd.read_csv(src, parse_dates=["Created at"])
    df["Revenue"] = df["Quantity"] * df["Unit Price"] - df["Discount"]
    df["Month"] = df["Created at"].dt.to_period("M").astype(str)
    monthly = df.groupby("Month").agg(Orders=("Order ID","nunique"), Units=("Quantity","sum"), Revenue=("Revenue","sum")).reset_index()
    monthly["Avg order value"] = (monthly["Revenue"]/monthly["Orders"]).round(2)
    top = df.groupby("Product").agg(Units=("Quantity","sum"), Revenue=("Revenue","sum")).sort_values("Revenue", ascending=False).reset_index()
    pivot = df.pivot_table(index="Product", columns="Month", values="Revenue", aggfunc="sum", fill_value=0).reset_index()
    with pd.ExcelWriter(out, engine="openpyxl") as xw:
        monthly.to_excel(xw, sheet_name="Monthly", index=False)
        top.to_excel(xw, sheet_name="Top products", index=False)
        pivot.to_excel(xw, sheet_name="Product x Month", index=False)
        for ws in xw.book.worksheets:
            for c in ws[1]:
                c.font = Font(bold=True, color="FFFFFF"); c.fill = PatternFill("solid", fgColor="2F5597")
            for i, col in enumerate(ws.columns, 1):
                ws.column_dimensions[get_column_letter(i)].width = max(len(str(c.value or "")) for c in col) + 3
            ws.freeze_panes = "A2"
        ws = xw.book["Monthly"]
        lc = LineChart(); lc.title = "Revenue by month"
        lc.add_data(Reference(ws, min_col=4, min_row=1, max_row=len(monthly)+1), titles_from_data=True)
        lc.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(monthly)+1))
        ws.add_chart(lc, "G2")
        ws = xw.book["Top products"]
        bc = BarChart(); bc.title = "Revenue by product"
        bc.add_data(Reference(ws, min_col=3, min_row=1, max_row=len(top)+1), titles_from_data=True)
        bc.set_categories(Reference(ws, min_col=1, min_row=2, max_row=len(top)+1))
        ws.add_chart(bc, "F2")
    print(f"Wrote {out}: {len(df)} order lines, total revenue {df['Revenue'].sum():,.2f}")

if __name__ == "__main__":
    build(sys.argv[1], sys.argv[2])
